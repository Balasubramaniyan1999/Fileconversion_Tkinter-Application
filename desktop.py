from tkinter import *
from tkinter import filedialog
import os
import csv
from fpdf import FPDF
import docx
from PIL import Image, ImageTk



#creating the application main window.
root = Tk()
root.minsize(1000, 1000)
root.title("FileConversion_Application")

image1 = Image.open(r"D:\Desktop_project\ilya-pavlov-hXrPSgGFpqQ-unsplash.jpg")
test = ImageTk.PhotoImage(image1)
label1 = Label(image=test)
label1.image = test

# Position image
label1.place(x=0, y=0)
#creating a text
l = Label(root, text="Text file conversion")
l.config(font=("Courier", 14))
l.pack()

data = ""
#browsing files

def browseFiles():
    global data
    filename = filedialog.askopenfilename(initialdir="/",title="Select a File",filetypes=(("Text files","*.txt*"),("all files", "*.*")))
    label_file_explorer.configure(text="File selected: " + filename)
    filename = open(filename)  # or tf = open(tf, 'r')
    data = filename.read()





label_file_explorer = Label(root,
                            text = "",
                            width = 100, height = 4,
                            fg = "blue")
#browse button
button_explore = Button(root,text = "Upload Files",fg="pink", bg="black",command = browseFiles)
button_explore.place(x=170, y=150)
label_file_explorer.place(x=140, y=130)


#Create a dropdown Menu
menu= StringVar()
menu.set("Select Extension")
drop= OptionMenu(root, menu,"json", "csv","pdf","docx","html")
drop.place(x=170, y=250)


def conversion():
    #print(data)
    if(menu.get()) == "csv":
        with open("converted.csv", 'w') as csvfile:
            csvwriter = csv.writer(csvfile)
            csvwriter.writerows(data)
    elif(menu.get()) == "json":
        with open("converted.json", "w") as outfile:
            outfile.write(data)
    elif(menu.get()) == "docx":
        my_doc = docx.Document()
        my_doc.add_paragraph(data)
        my_doc.save("converted.docx")
    elif(menu.get()) == "html":
        Func = open("converted.html", 'w')
        Func.write(data)
        Func.close()
    elif (menu.get()) == "pdf":
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=15)
        f = data
        #for x in f:
        #pdf.cell(w=0,h=0, txt=f,border=1, ln=2, align='C')
        pdf.cell(20, 10, f, 0, 2, 'C')
        pdf.output("converted.pdf")
    else:
        pass




b10 = Button(root, text="      convert     ", fg="pink", bg="black", command=conversion)
b10.place(x=170, y=350)






root.mainloop()


